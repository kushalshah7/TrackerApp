from __future__ import annotations

import os
from datetime import date, datetime
from io import BytesIO
from typing import Any

import openpyxl
import psycopg
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb

from .config import MODULES, validate_entry_data


SHEETS = {"weekly-review": ("Presales Review",), "weekly-meeting": ("Presales", "ISP Team")}
MEETING_HEADERS = MODULES["weekly-meeting"]["fields"]
MONTH_YEAR_FIELDS = {"Date of Opportunity (MM/YY)", "Expected Month"}
PRESALES_NAMES = (
    "Aditya Potdar",
    "Ankesh Singh",
    "Arun M",
    "Irshad",
    "Kalim Ansari",
    "Pawan Dubey",
    "Suraj Raskar",
    "Surender Kumar",
)
ACCOUNT_MANAGER_ALIASES = {
    "hrishi sir": "Hrishikesh Phadnis",
    "hrishikesh sir": "Hrishikesh Phadnis",
    "jai": "Jaidrath Maniyar",
    "krathika": "Kratika",
    "merlyn methew": "Merlyn Mathew",
    "mohit kapoor": "Mohit Kapoor",
    "moihit kapoor": "Mohit Kapoor",
    "noaman vohra": "Noaman Vohara",
    "navneet": "Navaneet",
}
EXCLUDED_ACCOUNT_MANAGER_NAMES = {"na", "team"}


class ConflictError(Exception):
    pass


def json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value


def normalize_month_year(value: Any) -> Any:
    if value in (None, ""):
        return value
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m")
    text = str(value).strip()
    for pattern in ("%Y-%m", "%Y-%m-%d", "%b-%y", "%B %Y"):
        try:
            return datetime.strptime(text, pattern).strftime("%Y-%m")
        except ValueError:
            continue
    raise ValueError("Month and year must use the YYYY-MM format")


def excel_month_year(value: Any) -> Any:
    normalized = normalize_month_year(value)
    if normalized in (None, ""):
        return normalized
    return datetime.strptime(normalized, "%Y-%m").date()


def canonical_full_names(values: list[str]) -> list[str]:
    """Collapse known aliases and unambiguous short/full-name variants."""
    variants: dict[str, dict[str, int]] = {}
    for value in values:
        name = value.strip()
        if (not name or name.casefold() in EXCLUDED_ACCOUNT_MANAGER_NAMES
                or not name.replace(" ", "").isalpha()):
            continue
        name = ACCOUNT_MANAGER_ALIASES.get(name.casefold(), name)
        key = name.casefold()
        spellings = variants.setdefault(key, {})
        spellings[name] = spellings.get(name, 0) + 1

    by_first_name: dict[str, list[str]] = {}
    for spellings in variants.values():
        display = max(spellings, key=lambda name: (spellings[name], sum(char.isupper() for char in name)))
        first_name = display.split()[0].casefold()
        by_first_name.setdefault(first_name, []).append(display)

    result = []
    for names in by_first_name.values():
        full_names = [name for name in names if len(name.split()) > 1]
        if full_names:
            result.extend(full_names)
        else:
            result.extend(names)
    return sorted(result, key=str.casefold)


class TrackerStore:
    def __init__(self, database_url: str | None = None):
        self.database_url = database_url or os.getenv("DATABASE_URL") or os.getenv("POSTGRES_URL")
        if not self.database_url:
            raise RuntimeError("DATABASE_URL is required")

    def connect(self):
        return psycopg.connect(self.database_url, row_factory=dict_row)

    def initialize(self):
        with self.connect() as connection:
            connection.execute("""
                create table if not exists tracker_entries (
                    id bigint generated always as identity primary key,
                    module text not null,
                    data jsonb not null,
                    created_at timestamptz not null default now(),
                    updated_at timestamptz not null default now()
                );
                create index if not exists tracker_entries_module_id_idx on tracker_entries(module, id desc);
                create table if not exists tracker_workbook_template (
                    singleton boolean primary key default true check (singleton),
                    content bytea not null,
                    sha256 text not null,
                    updated_at timestamptz not null default now()
                );
                alter table tracker_entries add column if not exists last_edited_at timestamptz;
                alter table tracker_entries enable row level security;
                alter table tracker_workbook_template enable row level security;
            """)

    def entries(self, module: str, limit: int | None = 5000, presales: str | None = None,
                before_id: int | None = None):
        query = ("select id, data, last_edited_at from tracker_entries "
                 "where module = %s and data->>'_deleted_at' is null")
        params: tuple[Any, ...] = (module,)
        if presales is not None:
            query += " and lower(data->>'Presales') = lower(%s)"
            params += (presales,)
        if before_id is not None:
            query += " and id < %s"
            params += (before_id,)
        query += " order by id desc"
        if limit is not None:
            query += " limit %s"
            params += (limit,)
        with self.connect() as connection:
            rows = connection.execute(query, params).fetchall()
        return [{**row["data"], "_row": row["id"],
                 "_last_edited_at": row["last_edited_at"].isoformat() if row["last_edited_at"] else None}
                for row in rows]

    def names(self, presales: str | None = None):
        query = ("select module, nullif(btrim(data->>'AM'), '') as am, "
                 "nullif(btrim(data->>'Client Manager'), '') as client_manager "
                 "from tracker_entries "
                 "where module in ('weekly-meeting', 'weekly-review') and data->>'_deleted_at' is null")
        params: tuple[Any, ...] = ()
        if presales is not None:
            query += " and lower(data->>'Presales') = lower(%s)"
            params = (presales,)
        with self.connect() as connection:
            rows = connection.execute(query, params).fetchall()
        account_managers = []
        for row in rows:
            field = "am" if row["module"] == "weekly-review" else "client_manager"
            value = row[field]
            if isinstance(value, str) and value.strip():
                account_managers.append(value)
        return {"am": canonical_full_names(account_managers),
                "presales": [presales] if presales else list(PRESALES_NAMES)}

    def _clean(self, module: str, data: dict[str, Any]):
        validate_entry_data(module, data)
        unknown = set(data) - set(MODULES[module]["fields"])
        if unknown:
            raise ValueError(f"Unsupported fields: {', '.join(sorted(unknown))}")
        clean = {key: normalize_month_year(value) if key in MONTH_YEAR_FIELDS else json_value(value)
                 for key, value in data.items()}
        if module == "weekly-meeting" and clean.get("Date"):
            clean["Month"] = date.fromisoformat(str(clean["Date"])[:10]).strftime("%B")
        for field in ("AM", "Client Manager"):
            value = clean.get(field)
            if not value:
                continue
            name = str(value).strip()
            if (name.casefold() in EXCLUDED_ACCOUNT_MANAGER_NAMES
                    or not name.replace(" ", "").isalpha()):
                raise ValueError(f"{field} must contain one person's name using letters and spaces")
            clean[field] = ACCOUNT_MANAGER_ALIASES.get(name.casefold(), name)
        return clean

    @staticmethod
    def _check_presales(data: dict[str, Any], previous: dict[str, Any] | None = None):
        name = data.get("Presales")
        if not name:
            return
        canonical = next((item for item in PRESALES_NAMES if item.casefold() == str(name).casefold()), None)
        if canonical:
            data["Presales"] = canonical
        elif previous is None or name != previous.get("Presales"):
            raise ValueError("Presales must be one of the eight team names")

    @staticmethod
    def _check_version(current: datetime | None, expected: datetime | None):
        if current != expected:
            raise ConflictError("This record changed since you opened it. Refresh the data and try again.")

    def add(self, module: str, data: dict[str, Any], allowed_presales: str | None = None):
        clean = self._clean(module, data)
        self._check_presales(clean)
        if allowed_presales is not None and clean.get("Presales") != allowed_presales:
            raise ValueError("You can only add entries under your Presales name")
        with self.connect() as connection:
            connection.execute("select pg_advisory_xact_lock(hashtext(%s))", (module,))
            serial = MODULES[module]["serial"]
            if serial:
                maximum = connection.execute(
                    "select coalesce(max((data->>%s)::int), 0) as maximum from tracker_entries "
                    "where module = %s and data->>%s ~ '^[0-9]+$' and coalesce(data->>'_source_sheet', 'Presales') = 'Presales'",
                    (serial, module, serial),
                ).fetchone()["maximum"]
                clean[serial] = maximum + 1
            row = connection.execute(
                "insert into tracker_entries(module, data, last_edited_at) values (%s, %s, now()) returning id",
                (module, Jsonb(clean)),
            ).fetchone()
        return {"row": row["id"], "message": "Entry added successfully"}

    def update(self, module: str, record_id: int, data: dict[str, Any], expected_last_edited_at: datetime | None,
               allowed_presales: str | None = None):
        clean = self._clean(module, data)
        with self.connect() as connection:
            old = connection.execute(
                "select data, last_edited_at from tracker_entries "
                "where id = %s and module = %s and data->>'_deleted_at' is null for update",
                (record_id, module),
            ).fetchone()
            if not old:
                raise ValueError("Entry was not found")
            if allowed_presales is not None and old["data"].get("Presales") != allowed_presales:
                raise ValueError("Entry was not found")
            self._check_version(old["last_edited_at"], expected_last_edited_at)
            self._check_presales(clean, old["data"])
            if allowed_presales is not None and clean.get("Presales") != allowed_presales:
                raise ValueError("You cannot change the Presales owner")
            for key in ("_source_row", "_source_sheet", MODULES[module]["serial"]):
                if key and key in old["data"]:
                    clean[key] = old["data"][key]
            connection.execute(
                "update tracker_entries set data = %s, updated_at = now(), last_edited_at = now() where id = %s",
                (Jsonb(clean), record_id),
            )
        return {"row": record_id, "message": "Entry updated successfully"}

    def delete(self, module: str, record_id: int, expected_last_edited_at: datetime | None,
               allowed_presales: str | None = None):
        with self.connect() as connection:
            current = connection.execute(
                "select data, last_edited_at from tracker_entries "
                "where module = %s and id = %s and data->>'_deleted_at' is null for update",
                (module, record_id),
            ).fetchone()
            if not current:
                raise ValueError("Entry was not found")
            if allowed_presales is not None and current["data"].get("Presales") != allowed_presales:
                raise ValueError("Entry was not found")
            self._check_version(current["last_edited_at"], expected_last_edited_at)
            connection.execute(
                "update tracker_entries set data = jsonb_set(data, '{_deleted_at}', to_jsonb(now()::text)), "
                "updated_at = now(), last_edited_at = now() where module = %s and id = %s",
                (module, record_id),
            )
        return {"row": record_id, "message": "Entry deleted successfully"}

    def workbook_bytes(self, presales: str | None = None):
        if presales is not None:
            workbook = openpyxl.Workbook()
            workbook.remove(workbook.active)
            for module in SHEETS:
                sheet = workbook.create_sheet(MODULES[module]["sheet"])
                fields = MODULES[module]["fields"]
                sheet.append(fields)
                for record in reversed(self.entries(module, limit=None, presales=presales)):
                    values = [excel_month_year(record.get(field)) if field in MONTH_YEAR_FIELDS
                              else record.get(field) for field in fields]
                    sheet.append(values)
                    for col, field in enumerate(fields, 1):
                        cell = sheet.cell(sheet.max_row, col)
                        if field in MONTH_YEAR_FIELDS:
                            cell.number_format = "mmm-yy"
                        elif isinstance(cell.value, str) and cell.value.startswith("="):
                            cell.data_type = "s"
            output = BytesIO()
            workbook.save(output)
            workbook.close()
            return output.getvalue()
        with self.connect() as connection:
            template = connection.execute(
                "select content from tracker_workbook_template where singleton = true"
            ).fetchone()
        if not template:
            raise ValueError("Workbook has not been imported")
        workbook = openpyxl.load_workbook(BytesIO(bytes(template["content"])))
        for module, names in SHEETS.items():
            records = list(reversed(self.entries(module, limit=None)))
            for sheet_name in names:
                sheet = workbook[sheet_name]
                fields = MODULES[module]["fields"]
                original_rows = [row for row in range(2, sheet.max_row + 1)
                                 if any(sheet.cell(row, col).value not in (None, "")
                                        for col in range(1, len(fields) + 1))]
                for row in original_rows:
                    for col in range(1, len(fields) + 1):
                        sheet.cell(row, col).value = None
                next_row = max(original_rows, default=1) + 1
                for record in records:
                    if record.get("_source_sheet", MODULES[module]["sheet"]) != sheet_name:
                        continue
                    target = record.get("_source_row") or next_row
                    if not record.get("_source_row"):
                        next_row += 1
                    for col, field in enumerate(fields, 1):
                        value = record.get(field)
                        cell = sheet.cell(target, col)
                        if field in MONTH_YEAR_FIELDS:
                            value = excel_month_year(value)
                            cell.number_format = "mmm-yy"
                        elif field == "Date" and isinstance(value, str):
                            try:
                                value = date.fromisoformat(value)
                            except ValueError:
                                pass
                        cell.value = value
                        if isinstance(value, str) and value.startswith("="):
                            cell.data_type = "s"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()
