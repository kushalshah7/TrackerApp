import os
import json
import time
import hashlib
from io import BytesIO
from types import SimpleNamespace

import jwt
import openpyxl
import pytest
from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
from fastapi import HTTPException
from fastapi.testclient import TestClient

os.environ.setdefault("DATABASE_URL", "postgresql://unused")

from app.auth.neon_auth import AccessSettings, AuthenticatedUser, TokenValidator, current_user, verified_identity  # noqa: E402
import app.auth.neon_auth as neon_auth  # noqa: E402
from app.main import app  # noqa: E402
import app.main as main  # noqa: E402
from app.models import EntryPayload  # noqa: E402
from app.tracker_store import TrackerStore  # noqa: E402
from app.tracker_store import PRESALES_NAMES  # noqa: E402


REVIEW = {"Region": "West", "Presales": "Pawan Dubey", "Customer": "Example",
          "Opportunity Details": "A meeting"}


class FakeStore:
    def __init__(self):
        self.calls = []

    def names(self, presales):
        self.calls.append(("names", presales))
        return {"am": [], "presales": [presales] if presales else []}

    def entries(self, module, limit, presales, before_id):
        self.calls.append(("entries", module, limit, before_id, presales))
        return []

    def add(self, module, data, presales):
        self.calls.append(("add", module, data, presales))
        return {"row": 1}

    def update(self, module, record_id, data, version, presales):
        self.calls.append(("update", module, record_id, version, presales))
        return {"row": record_id}

    def delete(self, module, record_id, version, presales):
        self.calls.append(("delete", module, record_id, version, presales))
        return {"row": record_id}

    def workbook_bytes(self, presales):
        self.calls.append(("download", presales))
        return b"workbook"


def test_every_data_route_requires_sign_in(monkeypatch):
    fake = FakeStore()
    monkeypatch.setattr(main, "store", fake)
    with TestClient(app) as client:
        requests = [
            ("GET", "/api/names", None),
            ("GET", "/api/me", None),
            ("GET", "/api/entries/weekly-review", None),
            ("POST", "/api/entries/weekly-review", {"data": REVIEW}),
            ("PATCH", "/api/entries/weekly-review/1", {"data": REVIEW, "expected_last_edited_at": None}),
            ("DELETE", "/api/entries/weekly-review/1", {"expected_last_edited_at": None}),
            ("GET", "/api/workbook/download", None),
            ("POST", "/api/auth/enroll", {"code": "a" * 32}),
        ]
        for method, path, body in requests:
            assert client.request(method, path, json=body).status_code == 401
    assert fake.calls == []


def test_presales_identity_is_applied_to_every_store_operation(monkeypatch):
    fake = FakeStore()
    monkeypatch.setattr(main, "store", fake)
    user = AuthenticatedUser("oid", "pawan.dubey@invecto.com", "Pawan Dubey", "Pawan Dubey")
    app.dependency_overrides[current_user] = lambda: user
    try:
        with TestClient(app) as client:
            assert client.get("/api/me").json()["presales"] == "Pawan Dubey"
            assert client.get("/api/names").status_code == 200
            assert client.get("/api/entries/weekly-review").status_code == 200
            assert client.post("/api/entries/weekly-review", json={"data": REVIEW}).status_code == 201
            assert client.patch("/api/entries/weekly-review/1", json={"data": REVIEW,
                                "expected_last_edited_at": None}).status_code == 200
            assert client.request("DELETE", "/api/entries/weekly-review/1",
                                  json={"expected_last_edited_at": None}).status_code == 200
            assert client.get("/api/workbook/download").content == b"workbook"
    finally:
        app.dependency_overrides.clear()
    assert [call[-1] for call in fake.calls] == ["Pawan Dubey"] * 6


def test_admin_identity_can_access_all_records(monkeypatch):
    fake = FakeStore()
    monkeypatch.setattr(main, "store", fake)
    app.dependency_overrides[current_user] = lambda: AuthenticatedUser(
        "oid", "kushal.shah@invecto.com", "Kushal Shah", None
    )
    try:
        with TestClient(app) as client:
            assert client.get("/api/me").json()["role"] == "admin"
            assert client.get("/api/entries/weekly-meeting").status_code == 200
            assert client.get("/api/workbook/download").status_code == 200
    finally:
        app.dependency_overrides.clear()
    assert fake.calls == [("entries", "weekly-meeting", 5000, None, None), ("download", None)]


def test_signed_in_but_uninvited_account_cannot_read_data(monkeypatch):
    fake = FakeStore()
    monkeypatch.setattr(main, "store", fake)
    monkeypatch.setattr(neon_auth, "is_enrolled", lambda _: False)
    app.dependency_overrides[verified_identity] = lambda: AuthenticatedUser(
        "user-id", "pawan.dubey@invecto.com", "Pawan Dubey", "Pawan Dubey"
    )
    try:
        with TestClient(app) as client:
            assert client.get("/api/entries/weekly-review").status_code == 403
            assert client.get("/api/workbook/download").status_code == 403
    finally:
        app.dependency_overrides.clear()
    assert fake.calls == []


def test_enrollment_requires_matching_invitation(monkeypatch):
    user = AuthenticatedUser("user-id", "pawan.dubey@invecto.com", "Pawan Dubey", "Pawan Dubey")
    redeemed = []
    monkeypatch.setattr(main, "redeem_invite", lambda identity, code: redeemed.append((identity, code)))
    app.dependency_overrides[verified_identity] = lambda: user
    try:
        with TestClient(app) as client:
            assert client.post("/api/auth/enroll", json={"code": "a" * 32}).status_code == 200
    finally:
        app.dependency_overrides.clear()
    assert redeemed == [(user, "a" * 32)]


def test_invitation_redemption_binds_code_to_email_and_auth_user(monkeypatch):
    user = AuthenticatedUser("neon-user-id", "pawan.dubey@invecto.com", "Pawan Dubey", "Pawan Dubey")

    class Connection:
        def __enter__(self):
            return self

        def __exit__(self, *_):
            pass

        def execute(self, query, params):
            self.query, self.params = query, params
            return self

        def fetchone(self):
            return ("pawan.dubey@invecto.com",)

    connection = Connection()
    monkeypatch.setenv("DATABASE_URL", "postgresql://unused")
    monkeypatch.setattr(neon_auth.psycopg, "connect", lambda _: connection)
    neon_auth.redeem_invite(user, "a" * 32)
    assert connection.params == (
        "neon-user-id", "pawan.dubey@invecto.com", hashlib.sha256(("a" * 32).encode()).hexdigest()
    )
    assert "redeemed_user_id is null" in connection.query
    assert "expires_at > now()" in connection.query


def test_signup_checks_invitation_before_account_creation(monkeypatch):
    checked = []
    monkeypatch.setattr(main, "check_invite", lambda email, code: checked.append((email, code)))
    with TestClient(app) as client:
        response = client.post("/api/auth/check-invite", json={
            "email": "pawan.dubey@invecto.com", "code": "a" * 32,
        })
    assert response.status_code == 200
    assert checked == [("pawan.dubey@invecto.com", "a" * 32)]


@pytest.fixture(scope="module")
def token_fixture():
    private_key = Ed25519PrivateKey.generate()
    settings = AccessSettings(
        auth_url="https://auth.example.test/neondb/auth",
        admin_emails=frozenset({"kushal.shah@invecto.com", "javed.khan@invecto.com"}),
        presales_by_email={"pawan.dubey@invecto.com": "Pawan Dubey"},
    )
    validator = TokenValidator(settings)
    validator.jwks = SimpleNamespace(
        get_signing_key_from_jwt=lambda _: SimpleNamespace(key=private_key.public_key())
    )
    now = int(time.time())
    claims = {"iss": settings.issuer, "aud": settings.issuer, "sub": "user-object-id",
              "email": "Pawan.Dubey@invecto.com", "name": "Pawan Dubey",
              "iat": now, "exp": now + 3600}

    def sign(changes=None):
        return jwt.encode({**claims, **(changes or {})}, private_key, algorithm="EdDSA")

    return validator, sign


def test_valid_tracker_account_maps_to_one_presales_identity(token_fixture):
    validator, sign = token_fixture
    user = validator.validate(sign())
    assert user.presales == "Pawan Dubey"
    assert user.email == "pawan.dubey@invecto.com"
    assert validator.validate(sign({"email": "javed.khan@invecto.com"})).is_admin


def test_access_configuration_requires_exactly_eight_distinct_presales(monkeypatch):
    values = {
        "NEON_AUTH_BASE_URL": "https://auth.example.test/neondb/auth",
        "TRACKER_ADMIN_EMAILS": "kushal.shah@invecto.com,javed.khan@invecto.com",
        "TRACKER_PRESALES_EMAIL_MAP": json.dumps({f"person{i}@invecto.com": name
                                                 for i, name in enumerate(PRESALES_NAMES)}),
    }
    for key, value in values.items():
        monkeypatch.setenv(key, value)
    settings = AccessSettings.from_env()
    assert len(settings.presales_by_email) == 8
    assert settings.issuer == "https://auth.example.test"
    assert len(settings.admin_emails) == 2

    monkeypatch.setenv("TRACKER_PRESALES_EMAIL_MAP", json.dumps({"one@invecto.com": PRESALES_NAMES[0]}))
    with pytest.raises(RuntimeError, match="each team member"):
        AccessSettings.from_env()


@pytest.mark.parametrize(
    ("changes", "status"),
    [({"email": "outsider@invecto.com"}, 403),
     ({"banned": True}, 403),
     ({"iss": "https://another-auth.example.test"}, 401),
     ({"aud": "another-api"}, 401),
     ({"exp": 1}, 401)],
)
def test_invalid_or_unlisted_tokens_are_rejected(token_fixture, changes, status):
    validator, sign = token_fixture
    with pytest.raises(HTTPException) as error:
        validator.validate(sign(changes))
    assert error.value.status_code == status


def test_presales_workbook_contains_only_own_rows_and_no_template_sheets(monkeypatch):
    store = TrackerStore("postgresql://unused")

    def own_rows(module, limit, presales):
        assert limit is None and presales == "Pawan Dubey"
        if module == "weekly-review":
            return [{"Region": "West", "Presales": "Pawan Dubey", "Customer": "=1+1",
                     "Date of Opportunity (MM/YY)": "2026-09"}]
        return []

    monkeypatch.setattr(store, "entries", own_rows)
    workbook = openpyxl.load_workbook(BytesIO(store.workbook_bytes("Pawan Dubey")))
    try:
        assert workbook.sheetnames == ["Presales Review", "Presales"]
        sheet = workbook["Presales Review"]
        assert sheet.max_row == 2
        assert sheet["D2"].value == "=1+1"
        assert sheet["D2"].data_type == "s"
        assert sheet["E2"].number_format == "mmm-yy"
        assert workbook["Presales"].max_row == 1
    finally:
        workbook.close()


def test_store_filters_reads_by_presales_identity(monkeypatch):
    class Connection:
        def __enter__(self):
            return self

        def __exit__(self, *_):
            pass

        def execute(self, query, params):
            self.query, self.params = query, params
            return self

        def fetchall(self):
            return []

    connection = Connection()
    store = TrackerStore("postgresql://unused")
    monkeypatch.setattr(store, "connect", lambda: connection)
    assert store.entries("weekly-review", 100, "Pawan Dubey", 20) == []
    assert "lower(data->>'Presales') = lower(%s)" in connection.query
    assert "id < %s" in connection.query
    assert connection.params == ("weekly-review", "Pawan Dubey", 20, 100)


def test_store_prevents_writes_to_another_presales_owner(monkeypatch):
    class Connection:
        def __init__(self):
            self.writes = []

        def __enter__(self):
            return self

        def __exit__(self, *_):
            pass

        def execute(self, query, params):
            if query.startswith("update") or query.startswith("insert"):
                self.writes.append(query)
            return self

        def fetchone(self):
            return {"data": {"Presales": "Ankesh Singh"}, "last_edited_at": None}

    connection = Connection()
    store = TrackerStore("postgresql://unused")
    monkeypatch.setattr(store, "connect", lambda: connection)
    with pytest.raises(ValueError, match="only add"):
        store.add("weekly-review", REVIEW, "Ankesh Singh")
    with pytest.raises(ValueError, match="not found"):
        store.update("weekly-review", 1, REVIEW, None, "Pawan Dubey")
    with pytest.raises(ValueError, match="not found"):
        store.delete("weekly-review", 1, None, "Pawan Dubey")
    assert connection.writes == []


def test_delete_marks_a_record_recoverable(monkeypatch):
    class Connection:
        def __enter__(self):
            return self

        def __exit__(self, *_):
            pass

        def execute(self, query, params):
            self.last_query = query
            return self

        def fetchone(self):
            return {"data": {"Presales": "Pawan Dubey"}, "last_edited_at": None}

    connection = Connection()
    store = TrackerStore("postgresql://unused")
    monkeypatch.setattr(store, "connect", lambda: connection)
    assert store.delete("weekly-review", 1, None, "Pawan Dubey")["row"] == 1
    assert connection.last_query.startswith("update tracker_entries")
    assert "_deleted_at" in connection.last_query


@pytest.mark.parametrize("invalid", ["NaN", "Infinity", "-1"])
def test_invalid_currency_amounts_are_rejected(invalid):
    with pytest.raises(ValueError):
        EntryPayload(data={"Value (₹)": invalid})
